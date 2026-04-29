"""
make_excel_template.py  (v006 - stations 전용)
===============================================
현재 config.py 값으로 input.xlsx 템플릿을 생성합니다.
처음 한 번만 실행하거나, 설정이 크게 바뀌었을 때 재실행합니다.

주의: input.xlsx가 이미 존재하면 config.py 임포트 시 자동 로드됩니다.
      올바른 템플릿을 생성하려면 기존 input.xlsx를 먼저 삭제하세요.

사용법:
  python make_excel_template.py             # input.xlsx 생성
  python make_excel_template.py output.xlsx # 파일명 직접 지정
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl 패키지가 필요합니다.\n설치: pip install openpyxl")

import config


# ──────────────────────────────────────────────────────────────
# 스타일 상수
# ──────────────────────────────────────────────────────────────
_HDR_FILL  = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
_HDR_FONT  = Font(bold=True, color='FFFFFF')
_KEY_FONT  = Font(bold=True)
_DESC_FONT = Font(color='595959', italic=True)
_CENTER    = Alignment(horizontal='center', vertical='center', wrap_text=False)
_LEFT      = Alignment(horizontal='left',   vertical='center')


def _hrow(ws, headers: list, row: int = 1):
    """헤더 행 작성 (파란 배경, 흰 볼드체)"""
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill      = _HDR_FILL
        c.font      = _HDR_FONT
        c.alignment = _CENTER


def _widths(ws, sizes: list):
    for i, w in enumerate(sizes, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ──────────────────────────────────────────────────────────────
# 시트 생성 함수
# ──────────────────────────────────────────────────────────────

def _sheet_basic(wb):
    ws = wb.create_sheet('기본설정')
    _hrow(ws, ['항목명', '값', '설명'])
    _widths(ws, [20, 38, 50])
    ws.freeze_panes = 'A2'

    rows = [
        ('HEC1_EXE',       config.HEC1_EXE,          'HEC-1 실행파일 전체 경로'),
        ('WORK_DIR',       config.WORK_DIR,           '계산 결과 기본 폴더 (결과는 하위폴더 project_name/method 에 저장됨)'),
        ('DT_MIN',         config.DT_MIN,             '계산 시간간격 (분)'),
        ('HUFF_QUARTILE',  config.HUFF_QUARTILE,      'Huff 분위 (1~4, 환경부 지침 채택=3)'),
        ('BASIN_AREA_KM2', config.BASIN_AREA_KM2,     '유역 대표면적 (km²) — 25.9 km² 이하 ARF=1.0'),
        ('ARF_REGION',     config.ARF_REGION,         '유역 본류 권역 : 한강 | 낙동강 | 금강 | 영산강'),
        ('PROJECT_NAME',   config.PROJECT['name'],    '프로젝트 이름 (결과 하위폴더명·HEC-1 ID카드에 사용)'),
    ]
    for r, (k, v, desc) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=k).font    = _KEY_FONT
        ws.cell(row=r, column=2, value=v)
        ws.cell(row=r, column=3, value=desc).font = _DESC_FONT


def _sheet_return_periods(wb):
    ws = wb.create_sheet('재현기간')
    _hrow(ws, ['재현기간(년)'])
    _widths(ws, [15])
    ws.freeze_panes = 'A2'
    for r, rp in enumerate(config.RETURN_PERIODS, start=2):
        ws.cell(row=r, column=1, value=rp).alignment = _CENTER


def _sheet_durations(wb):
    ws = wb.create_sheet('지속기간')
    _hrow(ws, ['지속기간(분)'])
    _widths(ws, [15])
    ws.freeze_panes = 'A2'
    for r, d in enumerate(config.DURATIONS, start=2):
        ws.cell(row=r, column=1, value=d).alignment = _CENTER


def _sheet_output_stations(wb):
    ws = wb.create_sheet('출력소')
    _hrow(ws, ['소명', 'VAR코드'])
    _widths(ws, [15, 12])
    ws.freeze_panes = 'A2'
    for r, st in enumerate(config.PROJECT['output_stations'], start=2):
        ws.cell(row=r, column=1, value=st['name'])
        ws.cell(row=r, column=2, value=st['var_code'])


def _basin_params(el: dict):
    """단위도 방법에 따라 (p1, p2) 반환 — 엑셀 tc/r 열에 저장."""
    method = el.get('method', 'clark').lower()
    params = el.get('params', {})
    if method == 'snyder':
        return params.get('tp'), params.get('cp')
    elif method == 'scs':
        return params.get('lag'), None
    else:
        return params.get('tc'), params.get('r')


def _sheet_network(wb):
    ws = wb.create_sheet('유역네트워크')
    headers = [
        '순서', 'type', 'comment', 'name',
        'area_km2', 'cn', 'method', 'tc', 'r', 'kk_line',
        'from_kk', 'k', 'x', 'n_reaches', 'n_inputs',
    ]
    _hrow(ws, headers)
    _widths(ws, [6, 12, 42, 12, 10, 8, 8, 8, 8, 14, 10, 8, 8, 12, 10])
    ws.freeze_panes = 'A2'

    type_notes = {
        'basin':      '← [Clark]tc/r  [Snyder]tc=tp,r=cp  [SCS]tc=lag  area_km2·cn·method 공통',
        'raw_basin':  '← [Clark]tc/r  [Snyder]tc=tp,r=cp  [SCS]tc=lag  kk_line 필수',
        'hydrograph': '← 외부수문곡선 시트에서 유량 배열 읽음 (QI 카드)',
        'reservoir':  '← 저수지_저류량 / 저수지_방류구 시트에서 데이터 읽음',
        'routing':    '← from_kk(KI 카드), k, x, n_reaches 사용',
        'combine':    '← n_inputs 사용, from_kk 지정 시 KI 카드 추가',
    }

    for seq, el in enumerate(config.PROJECT['network'], start=1):
        r     = seq + 1
        etype = el['type']
        ws.cell(row=r, column=1, value=seq).alignment  = _CENTER
        ws.cell(row=r, column=2, value=etype)
        ws.cell(row=r, column=3, value=el.get('comment', ''))
        ws.cell(row=r, column=4, value=el['name'])

        if etype in ('basin', 'raw_basin'):
            ws.cell(row=r, column=5,  value=el.get('area_km2'))
            ws.cell(row=r, column=6,  value=el.get('cn'))
            ws.cell(row=r, column=7,  value=el.get('method', 'clark'))
            p1, p2 = _basin_params(el)
            ws.cell(row=r, column=8,  value=p1)
            ws.cell(row=r, column=9,  value=p2)
            if etype == 'raw_basin':
                ws.cell(row=r, column=10, value=el.get('kk_line', ''))
        elif etype == 'routing':
            ws.cell(row=r, column=11, value=el.get('from_kk'))
            ws.cell(row=r, column=12, value=el.get('k'))
            ws.cell(row=r, column=13, value=el.get('x'))
            ws.cell(row=r, column=14, value=el.get('n_reaches', 1))
        elif etype == 'combine':
            ws.cell(row=r, column=11, value=el.get('from_kk'))
            ws.cell(row=r, column=15, value=el.get('n_inputs', 2))


def _sheet_storage_curves(wb):
    ws = wb.create_sheet('저수지_저류량')
    _hrow(ws, ['저수지명', '표고(m)', '저류량(천m³)'])
    _widths(ws, [12, 12, 14])
    ws.freeze_panes = 'A2'

    r = 2
    for el in config.PROJECT['network']:
        if el['type'] != 'reservoir':
            continue
        name = el['name']
        for el_v, st_v in zip(el.get('elevation', []), el.get('storage', [])):
            ws.cell(row=r, column=1, value=name)
            ws.cell(row=r, column=2, value=el_v)
            ws.cell(row=r, column=3, value=st_v)
            r += 1


def _sheet_stations(wb):
    """
    관측소_강우강도식: STATIONS 딕셔너리 → huff_cluster + 재현기간별 poly6 계수

    컬럼: 관측소명 | huff_cluster | 재현기간(년) | 식종류 | a0~a6
    """
    ws = wb.create_sheet('관측소_강우강도식')
    _hrow(ws, ['관측소명', 'huff_cluster', '재현기간(년)', '식종류',
               'a0', 'a1', 'a2', 'a3', 'a4', 'a5', 'a6'])
    _widths(ws, [12, 14, 14, 10] + [15] * 7)
    ws.freeze_panes = 'A2'

    r = 2
    for stn_name, stn_data in config.STATIONS.items():
        if stn_data is None:   # DUMMY는 행 없음 (자동 처리)
            continue
        cluster = stn_data.get('huff_cluster')
        # 정수 재현기간 키만 정렬 (huff_cluster 키는 제외)
        rp_keys = sorted(k for k in stn_data.keys() if isinstance(k, int))
        for rp in rp_keys:
            fc = stn_data[rp]
            ws.cell(row=r, column=1, value=stn_name).font       = _KEY_FONT
            ws.cell(row=r, column=2, value=cluster).alignment   = _CENTER
            ws.cell(row=r, column=3, value=rp).alignment        = _CENTER
            ws.cell(row=r, column=4, value=fc['formula'])
            for i in range(7):
                ws.cell(row=r, column=5 + i, value=fc['params'].get(f'a{i}', 0.0))
            r += 1


def _sheet_station_weights(wb):
    """유역_관측소가중치: basin별 station_weights → 유역명·관측소명·가중치"""
    ws = wb.create_sheet('유역_관측소가중치')
    _hrow(ws, ['유역명', '관측소명', '가중치'])
    _widths(ws, [12, 12, 12])
    ws.freeze_panes = 'A2'

    r = 2
    for el in config.PROJECT['network']:
        if el['type'] not in ('basin', 'raw_basin'):
            continue
        for stn_name, weight in el.get('station_weights', []):
            ws.cell(row=r, column=1, value=el['name'])
            ws.cell(row=r, column=2, value=stn_name)
            ws.cell(row=r, column=3, value=weight)
            r += 1


def _sheet_hydrographs(wb):
    """외부수문곡선: hydrograph 타입 요소의 QI 유량 배열"""
    ws = wb.create_sheet('외부수문곡선')
    _hrow(ws, ['지점명', '유량(m³/s)'])
    _widths(ws, [12, 14])
    ws.freeze_panes = 'A2'

    r = 2
    for el in config.PROJECT['network']:
        if el['type'] != 'hydrograph':
            continue
        for flow in el.get('ordinates', []):
            ws.cell(row=r, column=1, value=el['name'])
            ws.cell(row=r, column=2, value=flow)
            r += 1
    if r == 2:
        ws.cell(row=2, column=1, value='(지점명)').font = _DESC_FONT
        ws.cell(row=2, column=2, value=0.0).font        = _DESC_FONT


def _sheet_stage_discharge(wb):
    """
    저수지_방류량: 옵션 2 — 수위-방류량 테이블 직접 입력 (SQ 카드)

    이 시트에 데이터가 있는 저수지는 저수지_방류구(공식)보다 우선 적용됩니다.
    두 시트에 동시에 입력하지 마세요.
    """
    ws = wb.create_sheet('저수지_방류량')
    _hrow(ws, ['저수지명', '수위(m)', '방류량(m³/s)', '방류_KK'])
    _widths(ws, [12, 12, 14, 10])
    ws.freeze_panes = 'A2'

    # 안내 주석 행
    note = ws.cell(row=2, column=1,
                   value='← 수위-방류량 곡선이 있는 저수지만 여기에 입력 (저수지_방류구는 비워둘 것)')
    note.font = _DESC_FONT

    # 예시: config에 stage_discharge가 정의된 저수지가 있으면 출력
    r = 3
    for el in config.PROJECT['network']:
        if el['type'] != 'reservoir' or not el.get('stage_discharge'):
            continue
        kk = el.get('outflow_kk')
        for idx, (elev, flow) in enumerate(el['stage_discharge']):
            ws.cell(row=r, column=1, value=el['name'])
            ws.cell(row=r, column=2, value=elev)
            ws.cell(row=r, column=3, value=flow)
            if idx == 0:
                ws.cell(row=r, column=4, value=kk)
            r += 1


def _sheet_spillways(wb):
    """
    SS 카드: Q = 유량계수(C) × 마루길이(L) × (수위 - 마루표고)^유량지수(n)
    """
    ws = wb.create_sheet('저수지_방류구')
    _hrow(ws, ['저수지명', '마루표고(m)', '마루길이(m)', '유량계수(C)', '유량지수(n)', '방류_KK'])
    _widths(ws, [12, 12, 12, 12, 12, 10])
    ws.freeze_panes = 'A2'

    r = 2
    for el in config.PROJECT['network']:
        if el['type'] != 'reservoir':
            continue
        ws.cell(row=r, column=1, value=el['name'])
        for j, v in enumerate(el.get('spillway', []), start=2):
            ws.cell(row=r, column=j, value=v)
        ws.cell(row=r, column=6, value=el.get('outflow_kk'))
        r += 1


# ──────────────────────────────────────────────────────────────
# 메인
# ──────────────────────────────────────────────────────────────

def create_template(output_path: str = 'input.xlsx'):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 기본 Sheet1 제거

    _sheet_basic(wb)
    _sheet_return_periods(wb)
    _sheet_durations(wb)
    _sheet_stations(wb)
    _sheet_station_weights(wb)
    _sheet_output_stations(wb)
    _sheet_network(wb)
    _sheet_storage_curves(wb)
    _sheet_spillways(wb)
    _sheet_stage_discharge(wb)
    _sheet_hydrographs(wb)

    wb.save(output_path)
    abs_path = os.path.abspath(output_path)
    print(f"[완료] 템플릿 저장됨: {abs_path}")
    print()
    print("사용 방법:")
    print("  1. input.xlsx를 열어 원하는 값으로 수정하고 저장하세요.")
    print("  2. main.py를 실행하면 엑셀 값이 config.py보다 우선 적용됩니다.")
    print("  3. input.xlsx를 삭제하면 config.py 기본값으로 복원됩니다.")
    print()
    print("주의: 재생성 전 기존 input.xlsx를 삭제해야 올바른 템플릿이 생성됩니다.")


if __name__ == '__main__':
    if len(sys.argv) > 1:
        out = sys.argv[1]
    elif getattr(sys, 'frozen', False):
        # exe 실행 시: exe 파일 옆에 저장
        out = os.path.join(os.path.dirname(sys.executable), 'input.xlsx')
    else:
        out = 'input.xlsx'
    create_template(out)
