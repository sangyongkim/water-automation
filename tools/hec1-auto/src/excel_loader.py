"""
excel_loader.py  (v006 - stations 전용)
=========================================
input.xlsx에서 HEC-1 설정값을 읽어 config.py 변수를 덮어씁니다.

시트 구조:
  기본설정           : HEC1_EXE, WORK_DIR, DT_MIN 등 단순 설정값
  재현기간           : RETURN_PERIODS 목록
  지속기간           : DURATIONS 목록
  관측소_강우강도식  : STATIONS (관측소별 huff_cluster + 재현기간별 poly6 계수)
  유역_관측소가중치  : 유역별 관측소 티센 가중치
  출력소             : PROJECT['output_stations']
  유역네트워크       : PROJECT['network']
  저수지_저류량      : 저수지별 표고-저류량 곡선
  저수지_방류구      : 저수지별 여수로 및 방류 KK번호
  외부수문곡선       : 외부 수문곡선 유량 배열 (QI 카드)

사용법: config.py 끝부분에서 자동으로 호출됩니다.
직접 테스트: python excel_loader.py [파일경로]
"""

try:
    import openpyxl
except ImportError:
    raise ImportError(
        "openpyxl 패키지가 필요합니다.\n"
        "설치: pip install openpyxl"
    )


# ──────────────────────────────────────────────────────────────
# 개별 시트 파서
# ──────────────────────────────────────────────────────────────

def _read_basic(ws) -> dict:
    """기본설정 시트: 항목명-값 쌍 읽기"""
    raw = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        raw[str(row[0]).strip()] = row[1]

    type_map = {
        'HEC1_EXE':       str,
        'WORK_DIR':       str,
        'ARF_REGION':     str,
        'PROJECT_NAME':   str,
        'DT_MIN':         int,
        'HUFF_QUARTILE':  int,
        'BASIN_AREA_KM2': float,
    }
    cfg = {}
    for key, typ in type_map.items():
        if key in raw and raw[key] is not None:
            v = raw[key]
            cfg[key] = str(v).strip() if typ is str else typ(v)
    return cfg


def _read_column_list(ws, dtype=int) -> list:
    """헤더 1행, 데이터 2행~, 단일 컬럼 읽기"""
    return [
        dtype(row[0])
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True)
        if row[0] is not None and str(row[0]).strip()
    ]


def _read_stations(ws) -> dict:
    """
    관측소_강우강도식 시트 → STATIONS

    컬럼 순서:
      1: 관측소명
      2: huff_cluster
      3: 재현기간(년)
      4: 식종류
      5~11: a0~a6
    """
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        name = str(row[0]).strip()
        if name.upper() == 'DUMMY':
            result['DUMMY'] = None
            continue
        cluster = int(row[1]) if row[1] is not None else None
        rp      = int(row[2])
        formula = str(row[3]).strip() if row[3] else 'poly6'
        coefs   = {
            f'a{i}': float(row[4 + i])
            for i in range(7)
            if (4 + i) < len(row) and row[4 + i] is not None
        }
        if name not in result:
            result[name] = {}
        # huff_cluster은 관측소당 한 번만 저장
        if cluster is not None and 'huff_cluster' not in result[name]:
            result[name]['huff_cluster'] = cluster
        result[name][rp] = {'formula': formula, 'params': coefs}
    return result


def _read_station_weights(ws) -> dict:
    """유역_관측소가중치 시트 → {유역명: [(관측소명, 가중치), ...]}"""
    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        basin  = str(row[0]).strip()
        stn    = str(row[1]).strip()
        weight = float(row[2])
        result.setdefault(basin, []).append((stn, weight))
    return result


def _read_output_stations(ws) -> list:
    """출력소 시트 → PROJECT['output_stations']"""
    return [
        {'name': str(row[0]).strip(), 'var_code': str(row[1]).strip()}
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]


def _read_hydrographs(ws) -> dict:
    """외부수문곡선 시트 → {지점명: [유량(m³/s), ...]}"""
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        name = str(row[0]).strip()
        flow = float(row[1]) if row[1] is not None else 0.0
        data.setdefault(name, []).append(flow)
    return data


def _read_stage_discharge(ws) -> dict:
    """
    저수지_방류량 시트 → {저수지명: {'stage_discharge': [(수위, 방류량), ...], 'outflow_kk': int}}

    컬럼 순서: 저수지명 | 수위(m) | 방류량(m³/s) | 방류_KK
    방류_KK는 각 저수지의 첫 번째 행에만 입력하면 됩니다.
    """
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[1] is None or row[2] is None:
            continue
        name = str(row[0]).strip()
        elev = float(row[1])
        flow = float(row[2])
        kk   = int(row[3]) if len(row) > 3 and row[3] is not None else None
        entry = data.setdefault(name, {'stage_discharge': [], 'outflow_kk': None})
        entry['stage_discharge'].append((elev, flow))
        if kk is not None and entry['outflow_kk'] is None:
            entry['outflow_kk'] = kk
    return data


def _read_storage_curves(ws) -> dict:
    """저수지_저류량 시트 → {저수지명: {'elevation': [], 'storage': []}}"""
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        name  = str(row[0]).strip()
        entry = data.setdefault(name, {'elevation': [], 'storage': []})
        entry['elevation'].append(float(row[1]))
        entry['storage'].append(float(row[2]))
    return data


def _read_spillways(ws) -> dict:
    """저수지_방류구 시트 → {저수지명: {'spillway': [], 'outflow_kk': int}}"""
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        name = str(row[0]).strip()
        data[name] = {
            'spillway':   [float(v) for v in row[1:5] if v is not None],
            'outflow_kk': int(row[5]) if len(row) > 5 and row[5] is not None else None,
        }
    return data


def _read_network(ws, storage_data: dict, spillway_data: dict,
                  station_weights_data: dict = None,
                  hydro_data: dict = None,
                  stage_discharge_data: dict = None) -> list:
    """유역네트워크 시트 → PROJECT['network']"""
    headers = [c.value for c in ws[1]]
    col     = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}

    def g(row, name, default=None):
        i = col.get(name)
        if i is None or i >= len(row) or row[i] is None:
            return default
        return row[i]

    network = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        etype = str(g(row, 'type', '')).strip().lower()
        if not etype:
            continue
        name    = str(g(row, 'name',    '')).strip()
        comment = str(g(row, 'comment', '')).strip()

        if etype in ('basin', 'raw_basin'):
            method = str(g(row, 'method', 'clark')).strip().lower()
            p1     = float(g(row, 'tc', 0.0))
            p2     = float(g(row, 'r',  0.0))
            if method == 'snyder':
                params = {'tp': p1, 'cp': p2}
            elif method == 'scs':
                params = {'lag': p1}
            else:
                params = {'tc': p1, 'r': p2}
            elem = {
                'type': etype, 'comment': comment, 'name': name,
                'area_km2': float(g(row, 'area_km2', 0.0)),
                'cn':       float(g(row, 'cn',       0.0)),
                'method':   method,
                'params':   params,
            }
            kk = g(row, 'kk_line')
            if kk:
                elem['kk_line'] = str(kk).strip()
            if station_weights_data and name in station_weights_data:
                elem['station_weights'] = station_weights_data[name]

        elif etype == 'reservoir':
            elem = {'type': 'reservoir', 'comment': comment, 'name': name}
            if name in storage_data:
                elem.update(storage_data[name])
            # 옵션 2(수위-방류량 테이블)가 있으면 우선 적용, 없으면 옵션 1(공식) 사용
            if stage_discharge_data and name in stage_discharge_data:
                elem.update(stage_discharge_data[name])
            elif name in spillway_data:
                elem.update(spillway_data[name])

        elif etype == 'routing':
            fk = g(row, 'from_kk')
            elem = {
                'type': 'routing', 'comment': comment, 'name': name,
                'from_kk':   int(float(fk)) if fk else None,
                'k':         float(g(row, 'k',         0.0)),
                'x':         float(g(row, 'x',         0.0)),
                'n_reaches': int(float(g(row, 'n_reaches', 1))),
            }

        elif etype == 'combine':
            fk = g(row, 'from_kk')
            elem = {
                'type': 'combine', 'comment': comment, 'name': name,
                'from_kk':  int(float(fk)) if fk else None,
                'n_inputs': int(float(g(row, 'n_inputs', 2))),
            }

        elif etype == 'hydrograph':
            elem = {'type': 'hydrograph', 'comment': comment, 'name': name}
            if hydro_data and name in hydro_data:
                elem['ordinates'] = hydro_data[name]

        else:
            continue
        network.append(elem)
    return network


# ──────────────────────────────────────────────────────────────
# 메인 로더
# ──────────────────────────────────────────────────────────────

def load_config_from_excel(path: str) -> dict:
    """
    input.xlsx를 읽어 config.py 변수들을 덮어쓸 딕셔너리를 반환합니다.

    존재하지 않는 시트·빈 셀은 무시(기존 config.py 기본값 유지)됩니다.
    반환 딕셔너리에 '_PROJECT_OVERRIDES' 키가 있으면
    config.py의 PROJECT를 해당 값으로 부분 업데이트합니다.
    """
    wb  = openpyxl.load_workbook(path, data_only=True)
    sn  = set(wb.sheetnames)
    cfg = {}

    # ── 단순 설정값 ──────────────────────────────────────────
    if '기본설정' in sn:
        cfg.update(_read_basic(wb['기본설정']))

    # ── 리스트 ───────────────────────────────────────────────
    if '재현기간' in sn:
        lst = _read_column_list(wb['재현기간'])
        if lst:
            cfg['RETURN_PERIODS'] = lst
    if '지속기간' in sn:
        lst = _read_column_list(wb['지속기간'])
        if lst:
            cfg['DURATIONS'] = lst

    # ── 관측소 강우강도식 ─────────────────────────────────────
    if '관측소_강우강도식' in sn:
        d = _read_stations(wb['관측소_강우강도식'])
        if d:
            cfg['STATIONS'] = d

    # ── 저수지 보조 데이터 ────────────────────────────────────
    storage        = _read_storage_curves(wb['저수지_저류량']) if '저수지_저류량' in sn else {}
    spillway       = _read_spillways(wb['저수지_방류구'])      if '저수지_방류구' in sn else {}
    stage_discharge = _read_stage_discharge(wb['저수지_방류량']) if '저수지_방류량' in sn else {}

    # ── 유역별 관측소 가중치 ──────────────────────────────────
    stn_weights = _read_station_weights(wb['유역_관측소가중치']) \
                  if '유역_관측소가중치' in sn else {}

    # ── 외부 수문곡선 ─────────────────────────────────────────
    hydro_data = _read_hydrographs(wb['외부수문곡선']) \
                 if '외부수문곡선' in sn else {}

    # ── PROJECT 관련 ─────────────────────────────────────────
    proj_name = cfg.pop('PROJECT_NAME', None)
    huff_q    = cfg.get('HUFF_QUARTILE')
    output_st = _read_output_stations(wb['출력소'])     if '출력소'      in sn else None
    network   = _read_network(wb['유역네트워크'], storage, spillway,
                               stn_weights, hydro_data, stage_discharge) \
                                                        if '유역네트워크' in sn else None

    proj_overrides = {
        'name':            proj_name,
        'huff_quartile':   huff_q,
        'arf_region':      cfg.get('ARF_REGION'),
        'basin_area_km2':  cfg.get('BASIN_AREA_KM2'),
        'output_stations': output_st,
        'network':         network,
    }
    if any(v is not None for v in proj_overrides.values()):
        cfg['_PROJECT_OVERRIDES'] = proj_overrides

    return cfg


# ──────────────────────────────────────────────────────────────
# 직접 실행 시 내용 확인
# ──────────────────────────────────────────────────────────────
if __name__ == '__main__':
    import sys, pprint
    p = sys.argv[1] if len(sys.argv) > 1 else 'input.xlsx'
    pprint.pprint(load_config_from_excel(p))
